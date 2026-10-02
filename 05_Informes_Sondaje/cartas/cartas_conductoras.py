#!/usr/bin/env python3
"""
SOLUT - Generador de Cartas Conductoras (formato CC-SLAB-01, Rev. 01) para envío de muestras a laboratorio.

Uso:
  python cartas_conductoras.py carta_lonco_ITT-UDEC.json  salida.docx
  python cartas_conductoras.py carta_nueva_aldea_MSTD.json salida.docx

Dependencias: pip install python-docx
"""
import json
import os
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Inches, Pt, RGBColor

AQUI = os.path.dirname(os.path.abspath(__file__))
SERIF = "Bookman Old Style"
SANS = "Arial"
AZUL_ENCABEZADO = "B8CCE4"


# ---------------------------------------------------------------- utilidades
def fuente(run, nombre=SANS, tam=9, negrita=False, subrayado=False, cursiva=False, color=None, resaltado=False):
    run.font.name = nombre
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), nombre)
    run.font.size = Pt(tam)
    run.bold = negrita
    run.underline = subrayado
    run.italic = cursiva
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if resaltado:
        hl = OxmlElement("w:highlight")
        hl.set(qn("w:val"), "yellow")
        rpr.append(hl)


def parrafo(doc_o_cell, texto="", nombre=SERIF, tam=12, negrita=False, alinea=None, antes=0, despues=0,
            sangria_izq=None, sangria_prim=None, sangria_der=None, interlineado=None, **kw):
    p = doc_o_cell.add_paragraph()
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(antes), Pt(despues)
    if interlineado:
        pf.line_spacing = interlineado
    if alinea:
        p.alignment = alinea
    if sangria_izq is not None:
        pf.left_indent = Pt(sangria_izq)
    if sangria_prim is not None:
        pf.first_line_indent = Pt(sangria_prim)
    if sangria_der is not None:
        pf.right_indent = Pt(sangria_der)
    if texto:
        fuente(p.add_run(texto), nombre, tam, negrita, **kw)
    return p


def borde_celda(cell, **bordes):
    tcPr = cell._tc.get_or_add_tcPr()
    tb = tcPr.find(qn("w:tcBorders"))
    if tb is None:
        tb = OxmlElement("w:tcBorders")
        tcPr.append(tb)
    for lado in ("top", "left", "bottom", "right"):
        el = OxmlElement("w:" + lado)
        el.set(qn("w:val"), bordes.get(lado, "single"))
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        tb.append(el)


def sombreado(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), color)
    tcPr.append(sh)


def margenes_celda(cell, arriba=20, abajo=20, izq=60, der=60):
    tcPr = cell._tc.get_or_add_tcPr()
    m = OxmlElement("w:tcMar")
    for k, v in (("top", arriba), ("left", izq), ("bottom", abajo), ("right", der)):
        e = OxmlElement("w:" + k)
        e.set(qn("w:w"), str(v))
        e.set(qn("w:type"), "dxa")
        m.append(e)
    tcPr.append(m)


def ancho_celda(cell, pts):
    cell.width = Pt(pts)


def campo(par, instruccion, tam=10, negrita=True):
    """Inserta un campo (PAGE / NUMPAGES) en un párrafo."""
    for tipo, txt in (("begin", None), (None, instruccion), ("separate", None), (None, "1"), ("end", None)):
        r = par.add_run()
        fuente(r, SANS, tam, negrita)
        if tipo:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), tipo)
            r._element.append(fc)
        elif txt == instruccion:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = " %s " % instruccion
            r._element.append(it)
        else:
            t = OxmlElement("w:t")
            t.text = txt
            r._element.append(t)


def hipervinculo(par, url, texto, tam=10):
    rid = par.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                             is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rf = OxmlElement("w:rFonts")
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rf.set(qn(a), SANS)
    rpr.append(rf)
    c = OxmlElement("w:color")
    c.set(qn("w:val"), "0000FF")
    rpr.append(c)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(u)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(tam * 2)))
    rpr.append(sz)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = texto
    r.append(t)
    h.append(r)
    par._p.append(h)



ORDEN = {
    "tcPr": ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar",
             "textDirection", "tcFitText", "vAlign", "hideMark"],
    "trPr": ["cnfStyle", "divId", "gridBefore", "gridAfter", "wBefore", "wAfter", "cantSplit", "trHeight",
             "tblHeader", "tblCellSpacing", "jc", "hidden"],
    "pPr": ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr",
            "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap",
            "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid",
            "spacing", "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc", "textDirection",
            "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr", "pPrChange"],
    "rPr": ["rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike", "dstrike", "outline",
            "shadow", "emboss", "imprint", "noProof", "snapToGrid", "vanish", "webHidden", "color", "spacing", "w",
            "kern", "position", "sz", "szCs", "highlight", "u", "effect", "bdr", "shd", "fitText", "vertAlign",
            "rtl", "cs", "em", "lang", "eastAsianLayout", "specVanish", "oMath"],
}


def ordenar_esquema(raiz):
    """Reordena hijos de tcPr/trPr/pPr/rPr según el orden del esquema OOXML (Word es estricto)."""
    for nombre, orden in ORDEN.items():
        for el in raiz.iter(qn("w:" + nombre)):
            hijos = list(el)
            def clave(h):
                tag = h.tag.split("}")[1]
                return orden.index(tag) if tag in orden else len(orden)
            hijos_ord = sorted(hijos, key=clave)
            if hijos_ord != hijos:
                for h in hijos:
                    el.remove(h)
                for h in hijos_ord:
                    el.append(h)

# ---------------------------------------------------------------- encabezado y pie
def encabezado(sec, meta):
    hdr = sec.header
    hdr.is_linked_to_previous = False
    for p in list(hdr.paragraphs):
        p._element.getparent().remove(p._element)
    anchos = [135, 132, 136, 141]
    t = hdr.add_table(rows=3, cols=4, width=Pt(sum(anchos)))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, w in enumerate(anchos):
        for r in t.rows:
            ancho_celda(r.cells[i], w)
    logo = t.cell(0, 0).merge(t.cell(1, 0))
    p = logo.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(os.path.join(AQUI, "logo_solut.png"), width=Pt(100))
    celdas = {
        (0, 1): ("Revisión N°:", False, 9, WD_ALIGN_PARAGRAPH.LEFT),
        (0, 2): (meta["revision"], True, 10, WD_ALIGN_PARAGRAPH.CENTER),
        (1, 1): ("Inicio Aplicación:", False, 9, WD_ALIGN_PARAGRAPH.LEFT),
        (1, 2): (meta["inicio_aplicacion"], True, 10, WD_ALIGN_PARAGRAPH.CENTER),
    }
    for (r, c), (txt, neg, tam, al) in celdas.items():
        cp = t.cell(r, c).paragraphs[0]
        cp.alignment = al
        fuente(cp.add_run(txt), SANS, tam, neg)
    c = t.cell(0, 3).paragraphs[0]
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fuente(c.add_run("Código:"), SANS, 9)
    c.add_run().add_break(WD_BREAK.LINE)
    fuente(c.add_run(meta["codigo"]), SANS, 10, True)
    c = t.cell(1, 3).paragraphs[0]
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fuente(c.add_run("Página"), SANS, 9)
    c.add_run().add_break(WD_BREAK.LINE)
    campo(c, "PAGE")
    fuente(c.add_run(" de "), SANS, 10, True)
    campo(c, "NUMPAGES")
    tit = t.cell(2, 0).merge(t.cell(2, 3))
    tp = tit.paragraphs[0]
    tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fuente(tp.add_run("CARTA CONDUCTORA"), SANS, 12, True)
    for row in t.rows:
        for cell in row.cells:
            borde_celda(cell)
            margenes_celda(cell, 20, 20, 60, 60)
    for row in t.rows:
        for cell in row.cells:
            for par in cell.paragraphs:
                par.paragraph_format.space_after = Pt(0)
                par.paragraph_format.space_before = Pt(0)
    hdr.add_paragraph()  # espacio bajo la tabla


def pie(sec):
    ft = sec.footer
    ft.is_linked_to_previous = False
    p = ft.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pPr = p._p.get_or_add_pPr()
    bd = OxmlElement("w:pBdr")
    top = OxmlElement("w:top")
    top.set(qn("w:val"), "single")
    top.set(qn("w:sz"), "4")
    top.set(qn("w:space"), "8")
    top.set(qn("w:color"), "000000")
    bd.append(top)
    pPr.append(bd)
    fuente(p.add_run("MAIPU 1623 - CONCEPCION - "), SANS, 9)
    hipervinculo(p, "mailto:lrodriguez@solut.cl", "lrodriguez@solut.cl", 9)
    fuente(p.add_run(" - FONO +56936274496"), SANS, 9)


# ---------------------------------------------------------------- carta
def tabla_muestras(doc, grupos):
    anchos = [56, 115, 124, 78, 171]  # pt (total ~544)
    n_filas = 1 + sum(len(g["items"]) for g in grupos)
    t = doc.add_table(rows=n_filas, cols=5)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for r in t.rows:
        for i, w in enumerate(anchos):
            ancho_celda(r.cells[i], w)
    heads = [("Muestra\nN°", WD_ALIGN_PARAGRAPH.CENTER), ("Procedencia", WD_ALIGN_PARAGRAPH.CENTER),
             ("Coordenadas", WD_ALIGN_PARAGRAPH.CENTER), ("Fecha\nMuestreo", WD_ALIGN_PARAGRAPH.CENTER),
             ("Ensayos", WD_ALIGN_PARAGRAPH.CENTER)]
    for i, (h, al) in enumerate(heads):
        c = t.cell(0, i)
        p = c.paragraphs[0]
        p.alignment = al
        lineas = h.split("\n")
        for k, ln in enumerate(lineas):
            if k:
                p.add_run().add_break(WD_BREAK.LINE)
            fuente(p.add_run(ln), SANS, 9, True)
        sombreado(c, AZUL_ENCABEZADO)
    # fila de encabezado se repite en cada página
    trPr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    trPr.append(th)
    fila = 1
    for g in grupos:
        n = len(g["items"])
        for k, item in enumerate(g["items"]):
            c = t.cell(fila + k, 1)
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            fuente(p.add_run(item), SANS, 9)
        for col, txt in ((0, g["muestra"]), (2, g["coordenadas"]), (3, g["fecha"]), (4, g["ensayos"])):
            c = t.cell(fila, col)
            if n > 1:
                c = c.merge(t.cell(fila + n - 1, col))
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            c.vertical_alignment = 1  # centrado
            fuente(p.add_run(txt), SANS, 9)
        fila += n
    for row in t.rows:
        for cell in row.cells:
            borde_celda(cell)
            margenes_celda(cell, 15, 15, 60, 60)
            for par in cell.paragraphs:
                par.paragraph_format.space_after = Pt(0)
                par.paragraph_format.space_before = Pt(0)
        # no partir filas entre páginas
        trPr = row._tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit")
        trPr.append(cs)
    return t


def generar(datos, salida):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Pt(36)
    sec.top_margin = Pt(118)
    sec.bottom_margin = Pt(62)
    sec.header_distance = Pt(36)
    sec.footer_distance = Pt(24)
    st = doc.styles["Normal"]
    st.font.name = SERIF
    st.font.size = Pt(12)
    st.paragraph_format.space_after = Pt(0)
    st.paragraph_format.space_before = Pt(0)
    for p in list(doc.paragraphs):
        p._element.getparent().remove(p._element)

    meta = datos["encabezado"]
    encabezado(sec, meta)
    pie(sec)

    # fecha y número de carta
    p = parrafo(doc, alinea=WD_ALIGN_PARAGRAPH.RIGHT)
    fuente(p.add_run("CONCEPCION"), SERIF, 12, True)
    fuente(p.add_run(", %s.-" % datos["fecha_carta"]), SERIF, 12)
    cn = datos["numero_carta"]
    p = parrafo(doc, alinea=WD_ALIGN_PARAGRAPH.RIGHT)
    fuente(p.add_run("CNº "), SERIF, 12)
    fuente(p.add_run(cn), SERIF, 12, resaltado=("_" in cn))

    # destinatario
    d = datos["destinatario"]
    parrafo(doc, "Señores.", SERIF, 12, antes=10)
    parrafo(doc, d["empresa"], SANS, 9, True)
    for ln in d["lineas"]:
        parrafo(doc, ln, SERIF, 12, True, resaltado=("[" in ln))
    parrafo(doc, "Presente.", SERIF, 12, True)

    p = parrafo(doc, alinea=WD_ALIGN_PARAGRAPH.RIGHT, antes=16, sangria_der=40)
    fuente(p.add_run("Ref.: Solicitud de Ensayos"), SERIF, 12)
    parrafo(doc, "De nuestra consideración:", SERIF, 12, antes=14)

    n = sum(len(g["items"]) for g in datos["grupos"])
    pr = datos["proyecto"]
    parrafo(doc, "Mediante la presente, se hace ingreso de %d muestras de Suelos extraídas del Proyecto: “%s”." % (
        n, pr["nombre"]), SERIF, 11, antes=14, sangria_prim=15, alinea=WD_ALIGN_PARAGRAPH.JUSTIFY)
    for ln in (("Ubicación: %s" % pr["ubicacion"]), ("Mandante: %s" % pr["mandante"]),
               "Solicitante: Solut Ingeniería Spa.", "At: Sr. Felipe Matamala Andrades."):
        parrafo(doc, ln, SERIF, 11)
    parrafo(doc, "", SERIF, 11, despues=6)

    tabla_muestras(doc, datos["grupos"])

    # observaciones y cierre
    p = parrafo(doc, "Observaciones:", SANS, 10, True, antes=26, despues=8, subrayado=True)
    p.paragraph_format.keep_with_next = True
    for ob in datos["observaciones"]:
        q = parrafo(doc, "- " + ob, SANS, 9.5, despues=1)
        q.paragraph_format.keep_with_next = True
    q = parrafo(doc, "Sin otro particular,", SANS, 10, antes=30)
    q.paragraph_format.keep_with_next = True
    q = parrafo(doc, "Atte.", SANS, 10, antes=8)
    q.paragraph_format.keep_with_next = True
    f = datos["firma"]
    q = parrafo(doc, "", SANS, 10, antes=70)
    q.paragraph_format.keep_with_next = True
    for ln in f["lineas"]:
        q = parrafo(doc, ln, SANS, 10, alinea=WD_ALIGN_PARAGRAPH.CENTER, sangria_izq=150)
        q.paragraph_format.keep_with_next = True
    parrafo(doc, f["iniciales"], SANS, 10, antes=40)

    ordenar_esquema(doc.element)
    ordenar_esquema(sec.header._element)
    ordenar_esquema(sec.footer._element)
    doc.core_properties.title = "Carta Conductora - %s" % datos["destinatario"]["empresa"]
    doc.core_properties.author = "SOLUT Ingeniería SpA"
    doc.save(salida)
    return n


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    datos = json.load(open(sys.argv[1], encoding="utf8"))
    n = generar(datos, sys.argv[2])
    print("Carta generada:", sys.argv[2], "-", n, "muestras")
